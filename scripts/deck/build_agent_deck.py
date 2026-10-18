"""Build the sglang-maestro review deck: three agent loops, operator-to-model differentiation, one deep dive."""

import math
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from speaker_notes import notes_text

FONT = "Microsoft YaHei"
MONO = "Consolas"

# Theme "Graphite & Copper": graphite carries the deck, each loop owns one hue, amber appears only on dark pages.
G900 = "16202A"
G800 = "222E3A"
G700 = "35424F"
G500 = "66727F"
G400 = "8C96A1"
G300 = "C5CCD4"
G200 = "DDE2E7"
G100 = "EEF1F4"
G050 = "F6F7F9"
WHITE = "FFFFFF"
AMBER = "F2B233"

LOOPS = {
    "A": {"main": "3F63A8", "tint": "E8EDF7", "dark": "86A2DD"},
    "B": {"main": "1E8A6E", "tint": "E2F2EC", "dark": "55C3A0"},
    "C": {"main": "C0612C", "tint": "F8EADF", "dark": "E99A66"},
}
OK, WARN, BAD = "1E8A6E", "C98A12", "B4382F"
SHARE_GEMM, SHARE_ATTN, SHARE_OTHER = "7D8894", LOOPS["C"]["main"], "D3D8DE"

SW, SH = 13.333, 7.5
MX = 0.6
CW = SW - 2 * MX
TOTAL = 9
Y_FOOT = 6.98

MIDDLE, TOP = MSO_ANCHOR.MIDDLE, MSO_ANCHOR.TOP
CENTER, LEFT, RIGHT = PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.RIGHT
DASH = MSO_LINE_DASH_STYLE.DASH


def rgb(value):
    return RGBColor.from_string(value)


def drop_theme_style(shape):
    # python-pptx shapes reference theme styles whose effect styles carry an outer shadow.
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def style_run(run, size, color, bold=False, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        element = rpr.find(qn(tag))
        if element is None:
            element = etree.SubElement(rpr, qn(tag))
        element.set("typeface", FONT)


def P(text, size, color, bold=False, font=FONT, **options):
    spec = {"runs": [(text, size, color, bold, font)]}
    spec.update(options)
    return spec


def runs(parts, **options):
    """parts: (text, size, color[, bold[, font]]) tuples rendered as one paragraph."""
    spec = {"runs": list(parts)}
    spec.update(options)
    return spec


def setup_frame(frame, anchor, inset):
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = anchor
    left, top, right, bottom = inset
    frame.margin_left = Inches(left)
    frame.margin_top = Inches(top)
    frame.margin_right = Inches(right)
    frame.margin_bottom = Inches(bottom)


def fill_paragraphs(frame, paragraphs):
    for index, spec in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = spec.get("align", LEFT)
        paragraph.space_before = Pt(spec.get("before", 0))
        paragraph.space_after = Pt(spec.get("after", 0))
        paragraph.line_spacing = spec.get("spacing", 1.0)
        ppr = paragraph._p.get_or_add_pPr()
        ppr.set("latinLnBrk", "0")
        bullet = spec.get("bullet")
        if bullet:
            indent = int(Inches(0.16))
            ppr.set("marL", str(indent))
            ppr.set("indent", str(-indent))
            color = etree.SubElement(ppr, qn("a:buClr"))
            etree.SubElement(color, qn("a:srgbClr")).set("val", bullet)
            etree.SubElement(ppr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(ppr, qn("a:buChar")).set("char", "•")
        for text, size, color, *rest in spec["runs"]:
            run = paragraph.add_run()
            run.text = text
            bold = rest[0] if rest else False
            style_run(run, size, color, bold, rest[1] if len(rest) > 1 else FONT)


def add_box(slide, x, y, w, h, fill=None, line=None, line_w=0.75, dash=None,
            kind=MSO_SHAPE.RECTANGLE, radius=None):
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    drop_theme_style(shape)
    if radius is not None:
        shape.adjustments[0] = radius
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_w)
        if dash is not None:
            shape.line.dash_style = dash
    else:
        shape.line.fill.background()
    return shape


def add_text(slide, x, y, w, h, paragraphs, anchor=TOP, inset=(0.0, 0.0, 0.0, 0.0)):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    setup_frame(box.text_frame, anchor, inset)
    fill_paragraphs(box.text_frame, paragraphs)
    return box


def boxed_text(shape, paragraphs, anchor=MIDDLE, inset=(0.1, 0.04, 0.1, 0.04)):
    setup_frame(shape.text_frame, anchor, inset)
    fill_paragraphs(shape.text_frame, paragraphs)
    return shape


def add_arrowheads(line, head=False, tail=True):
    # a:headEnd and a:tailEnd must follow the fill and dash children of a:ln, so style the line first.
    ln = line._get_or_add_ln()
    for wanted, tag in ((head, "a:headEnd"), (tail, "a:tailEnd")):
        if wanted:
            end = etree.SubElement(ln, qn(tag))
            end.set("type", "triangle")
            end.set("w", "med")
            end.set("len", "med")


def add_line(slide, x1, y1, x2, y2, color, width=1.5, head=False, tail=False, dash=None):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    drop_theme_style(connector)
    connector.line.color.rgb = rgb(color)
    connector.line.width = Pt(width)
    if dash is not None:
        connector.line.dash_style = dash
    add_arrowheads(connector.line, head, tail)
    return connector


def add_arrow(slide, x1, y1, x2, y2, color, width=1.5, both=False):
    return add_line(slide, x1, y1, x2, y2, color, width, head=both, tail=True)


def add_path_arrow(slide, points, color, width=1.5, dash=None):
    """An open polyline through points (inches) with an arrowhead at the last point."""
    (x0, y0), rest = points[0], points[1:]
    builder = slide.shapes.build_freeform(Inches(x0), Inches(y0))
    builder.add_line_segments([(Inches(x), Inches(y)) for x, y in rest], close=False)
    shape = builder.convert_to_shape()
    drop_theme_style(shape)
    shape.fill.background()
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if dash is not None:
        shape.line.dash_style = dash
    add_arrowheads(shape.line)
    return shape


def set_background(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


# ---------------------------------------------------------------------- motif: the loop ring
def ring(slide, cx, cy, d, color, width=2.0):
    """A circular arrow: an arc running clockwise from 1 o'clock round to just before 12 o'clock."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ARC, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d)
    )
    drop_theme_style(shape)
    # Arc angles are clockwise from 3 o'clock in 60000ths of a degree; python-pptx scales them by 1e-5.
    shape.adjustments[0] = 300 * 0.6
    shape.adjustments[1] = 250 * 0.6
    shape.fill.background()
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    add_arrowheads(shape.line)
    return shape


def ring_badge(slide, cx, cy, d, color, letter, size, width=1.75, letter_color=None):
    ring(slide, cx, cy, d, color, width)
    add_text(slide, cx - d / 2, cy - d / 2, d, d,
             [P(letter, size, letter_color or color, True, "Arial", align=CENTER)], anchor=MIDDLE)


# ---------------------------------------------------------------------- page furniture
def header(slide, eyebrow, title, subtitle=None, loop=None):
    color = LOOPS[loop]["main"] if loop else G500
    x = MX
    if loop:
        ring_badge(slide, MX + 0.17, 0.58, 0.34, color, loop, 10, width=1.5)
        x = MX + 0.46
    add_text(slide, x, 0.42, 9.0, 0.32, [P(eyebrow, 11, color, True, MONO)], anchor=MIDDLE)
    add_text(slide, MX, 0.8, CW, 0.58, [P(title, 25, G900, True)], anchor=MIDDLE)
    if subtitle:
        add_text(slide, MX, 1.38, CW, 0.36, [P(subtitle, 13, G500)], anchor=MIDDLE)


def footer(slide, page, note=None, dark=False):
    color = G400 if dark else G500
    add_text(slide, SW - MX - 1.0, Y_FOOT, 1.0, 0.28,
             [P(f"{page:02d} / {TOTAL:02d}", 10, color, False, MONO, align=RIGHT)], anchor=MIDDLE)
    if note:
        add_text(slide, MX, Y_FOOT, CW - 1.3, 0.28, [P(note, 10, color)], anchor=MIDDLE)


def chip_row(slide, x, y, w, h, labels, color, gap=0.3, size=11, fill=WHITE, planned=()):
    chip_w = (w - gap * (len(labels) - 1)) / len(labels)
    for index, label in enumerate(labels):
        left = x + index * (chip_w + gap)
        later = index in planned
        chip = add_box(slide, left, y, chip_w, h, fill=fill, line=G400 if later else color, line_w=1.0,
                       dash=DASH if later else None, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        boxed_text(chip, [P(label, size, G500 if later else G900, True, align=CENTER)],
                   inset=(0.03, 0, 0.03, 0))
        if index:
            add_arrow(slide, left - gap + 0.06, y + h / 2, left - 0.06, y + h / 2, color, 1.25)


def number_dot(slide, x, y, d, color, text, size=11, dashed=False):
    dot = add_box(slide, x, y, d, d, fill=None if dashed else color, line=color if dashed else None,
                  line_w=1.25, dash=DASH if dashed else None, kind=MSO_SHAPE.OVAL)
    boxed_text(dot, [P(text, size, color if dashed else WHITE, True, "Arial", align=CENTER)],
               inset=(0, 0, 0, 0))
    return dot


def status_chip(slide, x, y, w, h, color, label):
    add_box(slide, x, y, w, h, fill=G050, line=G200, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    d = 0.12
    add_box(slide, x + 0.11, y + (h - d) / 2, d, d, fill=color, kind=MSO_SHAPE.OVAL)
    add_text(slide, x + 0.3, y, w - 0.32, h, [P(label, 10.5, G700)], anchor=MIDDLE)


def tag(slide, x, y, w, h, label, planned):
    shape = add_box(slide, x, y, w, h, fill=None if planned else LOOPS["B"]["tint"],
                    line=G400 if planned else None, line_w=0.75, dash=DASH if planned else None,
                    kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    boxed_text(shape, [P(label, 9.5, G500 if planned else LOOPS["B"]["main"], True, align=CENTER)],
               inset=(0.02, 0, 0.02, 0))


# ---------------------------------------------------------------------- 1 cover
def cover_slide(slide):
    set_background(slide, G900)
    add_text(slide, 0.8, 0.95, 7.0, 0.34, [P("SGLANG-MAESTRO · 方案评审", 12, AMBER, True, MONO)],
             anchor=MIDDLE)
    add_text(slide, 0.8, 1.5, 7.3, 2.0, [
        P("从算子级到模型级", 44, WHITE, True, after=8),
        P("SGLang 适配调优 Agent", 30, "D5DCE3", True),
    ])
    add_text(slide, 0.8, 3.8, 6.5, 1.0, [
        P("面向 M100 芯片的单机部署，用三个 Agent Loop 分别助力：", 15, G300, spacing=1.2),
        P("框架迭代、模型适配与评估优化、算子绑定评估与优化", 15, G300, spacing=1.2),
    ])
    rings = (
        ("A", 9.95, 2.45, ["框架迭代"]),
        ("B", 8.95, 4.2, ["模型适配", "评估优化"]),
        ("C", 10.95, 4.2, ["算子绑定", "评估优化"]),
    )
    d = 2.15
    for letter, cx, cy, lines in rings:
        color = LOOPS[letter]["dark"]
        ring(slide, cx, cy, d, color, width=3.0)
        add_text(slide, cx - 0.8, cy - 0.62, 1.6, 1.24,
                 [P(letter, 28, color, True, "Arial", align=CENTER, after=2)]
                 + [P(line, 12.5, WHITE, True, align=CENTER) for line in lines], anchor=MIDDLE)
    add_text(slide, 8.2, 5.55, 3.5, 0.3, [P("三个 Loop 共用一套证据底座", 11, G400, align=CENTER)],
             anchor=MIDDLE)
    add_text(slide, 0.8, 6.4, 9.0, 0.3, [P(
        "离线可跑 · 结论可追溯到 SGLang 源码与实测数据 · 知识快照：SGLang 72730d6（2026-10-14）",
        10.5, G400, False, MONO)], anchor=MIDDLE)


# ---------------------------------------------------------------------- 2 overview
def loop_block(slide, x, y, w, h, loop, title, question, steps, planned=()):
    color = LOOPS[loop]["main"]
    add_box(slide, x, y, w, h, fill=LOOPS[loop]["tint"])
    ring_badge(slide, x + 0.46, y + 0.5, 0.56, color, loop, 15, width=2.0)
    add_text(slide, x + 0.9, y + 0.16, w - 1.0, 0.36, [P(title, 14.5, color, True)], anchor=MIDDLE)
    add_text(slide, x + 0.9, y + 0.56, w - 0.98, 0.5, [P(question, 11.5, G700, spacing=1.1)])
    chip_row(slide, x + 0.22, y + h - 0.66, w - 0.44, 0.44, steps, color, gap=0.28, size=10.5,
             planned=planned)


def overview_slide(slide):
    header(slide, "全局 · 三个 AGENT LOOP",
           "三个 Agent Loop 分层咬合：框架定规则，模型定目标，算子定成本",
           "模型适配与算子优化互为输入：模型告诉算子先做什么，算子告诉模型实际要花多少")
    a = LOOPS["A"]["main"]
    ay, ah = 1.95, 1.1
    add_box(slide, MX, ay, CW, ah, fill=LOOPS["A"]["tint"])
    ring_badge(slide, MX + 0.5, ay + ah / 2, 0.62, a, "A", 16, width=2.0)
    add_text(slide, MX + 1.0, ay + 0.14, 4.0, 0.36, [P("Loop A  助力框架迭代", 14.5, a, True)],
             anchor=MIDDLE)
    add_text(slide, MX + 1.0, ay + 0.52, 3.75, 0.5,
             [P(line, 11.5, G700, spacing=1.1) for line in ("这一版 SGLang 在 M100 上能开哪些特性、", "插件要补什么")])
    chip_row(slide, 5.35, ay + (ah - 0.46) / 2, CW + MX - 5.35 - 0.25, 0.46,
             ["源码能力快照", "特性可行性", "收益估算", "合并 launch"], a, gap=0.32, size=11)

    by, bh, bw = 3.66, 1.92, 5.0
    left_x, right_x = MX, SW - MX - bw
    loop_block(slide, left_x, by, bw, bh, "B", "Loop B  助力模型适配与评估优化",
               "新模型能不能上、怎么配、单机 SLO 内能扛多少", ["适配分析", "精度布局", "仿真评估", "压测校准"])
    loop_block(slide, right_x, by, bw, bh, "C", "Loop C  助力算子绑定评估与优化",
               "每个算子实际跑哪条实现、慢在哪、先优化谁", ["绑定解析", "实测对比", "优化手段", "A/B 复测"],
               planned=(3,))

    for cx in (left_x + bw / 2, right_x + bw / 2):
        add_arrow(slide, cx, ay + ah + 0.05, cx, by - 0.05, a, 1.5, both=True)
    add_text(slide, left_x + bw / 2 + 0.25, ay + ah + 0.04, right_x - left_x - 0.5, 0.5, [
        P("↓ 提供规则：启动参数、特性约束、平台 hook、算子分发规则", 10.5, G700, align=CENTER),
        P("↑ 回传缺口：OOT 平台不支持的子系统、待实现的 hook", 10.5, G700, align=CENTER),
    ], anchor=MIDDLE)

    gap_x0, gap_x1 = left_x + bw, right_x
    b_color, c_color = LOOPS["B"]["main"], LOOPS["C"]["main"]
    y_fwd, y_back = by + 0.66, by + 1.3
    add_arrow(slide, gap_x0 + 0.1, y_fwd, gap_x1 - 0.1, y_fwd, b_color, 1.75)
    add_text(slide, gap_x0 + 0.02, y_fwd - 0.5, gap_x1 - gap_x0 - 0.04, 0.46, [
        P("算子缺口 · 热点", 10.5, G900, True, align=CENTER),
        P("按占 step 比例排序", 10, G500, align=CENTER),
    ], anchor=MIDDLE)
    add_arrow(slide, gap_x1 - 0.1, y_back, gap_x0 + 0.1, y_back, c_color, 1.75)
    add_text(slide, gap_x0 + 0.02, y_back + 0.05, gap_x1 - gap_x0 - 0.04, 0.46, [
        P("实测耗时", 10.5, G900, True, align=CENTER),
        P("校准成本模型、重新规划", 10, G500, align=CENTER),
    ], anchor=MIDDLE)

    fy, fh = 5.86, 0.86
    add_box(slide, MX, fy, CW, fh, fill=G100)
    add_text(slide, MX + 0.25, fy, 1.6, fh, [P("共用证据底座", 13, G900, True)], anchor=MIDDLE)
    items = (
        ("SGLang 源码静态分析", "不导入 SGLang，每条事实带源码位置"),
        ("算子级 step 成本模型", "roofline + launch 间隙 + α-β 通信"),
        ("实测校准", "server 日志与压测结果拟合系数"),
    )
    item_x, item_w = 2.4, 3.42
    for index, (name, detail) in enumerate(items):
        x = item_x + index * item_w
        if index:
            add_line(slide, x - 0.12, fy + 0.18, x - 0.12, fy + fh - 0.18, G300, 0.75)
        add_text(slide, x, fy, item_w - 0.3, fh, [
            P(name, 12, G900, True, after=2),
            P(detail, 10.5, G500),
        ], anchor=MIDDLE)
    footer(slide, 2, "虚线 = 规划中　·　与原场景的对应：A ≈ 场景 3（框架特性）　B ≈ 场景 1 + 5（适配、部署）　"
                     "C ≈ 场景 2（算子性能）")


# ---------------------------------------------------------------------- 3 loop A
def loop_a_slide(slide):
    a, tint = LOOPS["A"]["main"], LOOPS["A"]["tint"]
    header(slide, "LOOP A · 助力框架迭代",
           "SGLang 每升一版，Agent 重算 M100 能开哪些特性、怎么调",
           "知识来自源码静态分析，而不是人工维护的经验清单：换版本重跑一次同步，结论跟着版本走", loop="A")
    py, ph = 1.98, 1.5
    trigger = add_box(slide, MX, py, 1.7, ph, fill=G100)
    boxed_text(trigger, [
        P("触发", 11, G500, True, after=4),
        P("SGLang 新版本", 12, G900, True, after=2),
        P("新特性", 12, G900, True, after=2),
        P("M100 插件更新", 12, G900, True),
    ], inset=(0.16, 0.06, 0.08, 0.06))
    stages = (
        ("① 同步知识", ("静态分析源码", "生成能力快照", "不导入 SGLang", "无需 GPU")),
        ("② 判可行", ("复用 SGLang", "自身的校验规则", "核对插件 hook", "与平台能力")),
        ("③ 算收益", ("成本模型估算：", "KV 容量", "decode 吞吐", "TTFT · TPOT")),
        ("④ 出配置", ("可共存的特性", "合并成一条 launch", "会被拒绝的组合", "提前剔除")),
    )
    sx0, sw, sgap = 2.62, 1.76, 0.32
    centers = []
    for index, (name, lines) in enumerate(stages):
        x = sx0 + index * (sw + sgap)
        centers.append(x + sw / 2)
        box = add_box(slide, x, py, sw, ph, fill=WHITE, line=a, line_w=1.0)
        boxed_text(box, [P(name, 13, a, True, after=6)] + [P(line, 11, G700, after=2) for line in lines],
                   anchor=TOP, inset=(0.14, 0.14, 0.08, 0.06))
        add_arrow(slide, x - sgap + 0.05, py + ph / 2, x - 0.05, py + ph / 2, a, 1.5)
    out_x = sx0 + 4 * sw + 3 * sgap + sgap
    add_arrow(slide, out_x - sgap + 0.05, py + ph / 2, out_x - 0.05, py + ph / 2, a, 1.5)
    output = add_box(slide, out_x, py, SW - MX - out_x, ph, fill=tint)
    boxed_text(output, [
        P("产出", 11, a, True, after=4),
        P("合并后的 launch", 11.5, G900, True, after=3),
        P("插件 hook 清单", 11.5, G900, True, after=3),
        P("框架缺口清单", 11.5, G900, True),
    ], inset=(0.14, 0.06, 0.08, 0.06))

    ry = py + ph + 0.28
    add_path_arrow(slide, [(centers[3], py + ph), (centers[3], ry), (centers[0], ry), (centers[0], py + ph)],
                   a, 1.25)
    add_text(slide, centers[0] + 0.2, ry + 0.04, 4.55, 0.3,
             [P("压测校准后更新结论；SGLang 升级时整条链重跑", 10.5, G500)], anchor=MIDDLE)
    planned = add_box(slide, centers[0] + 4.85, ry + 0.06, 3.5, 0.28, line=G400, dash=DASH,
                      kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    boxed_text(planned, [P("规划中：两版快照自动 diff，只重算受影响的结论", 9.5, G500, align=CENTER)],
               inset=(0.04, 0, 0.04, 0))

    ey, eh = 4.3, 2.45
    lw = 6.0
    add_box(slide, MX, ey, lw, eh, fill=G050, line=G200)
    add_text(slide, MX + 0.2, ey + 0.12, lw - 0.4, 0.36,
             [runs([("能力快照里有什么", 13, G900, True), ("　示例：SGLang 72730d6", 10.5, G500, False)])],
             anchor=MIDDLE)
    tiles = (("521", "启动参数"), ("284", "模型架构"), ("37", "平台 hook"), ("56", "算子分发类"), ("9", "OOT 守卫"))
    tile_w, tile_gap = 1.04, 0.1
    for index, (value, label) in enumerate(tiles):
        x = MX + 0.2 + index * (tile_w + tile_gap)
        add_box(slide, x, ey + 0.6, tile_w, 0.98, fill=WHITE, line=G200)
        add_text(slide, x, ey + 0.66, tile_w, 0.52, [P(value, 24, a, True, "Arial", align=CENTER)],
                 anchor=MIDDLE)
        add_text(slide, x, ey + 1.16, tile_w, 0.32, [P(label, 10.5, G700, align=CENTER)], anchor=MIDDLE)
    add_text(slide, MX + 0.2, ey + 1.72, lw - 0.4, 0.66, [P(
        "每条事实带源码位置（文件:行号）。特性顾问覆盖 DP attention、prefill / decode CP、LayerNorm SP、"
        "DeepEP、EPLB、TBO / SBO、HiCache、MTP / EAGLE / NGRAM 等", 10.5, G500, spacing=1.1)])

    rx = MX + lw + 0.33
    rw = SW - MX - rx
    add_box(slide, rx, ey, rw, eh, fill=tint)
    add_text(slide, rx + 0.22, ey + 0.12, rw - 0.4, 0.36,
             [P("例：Agent 交给框架迭代的“缺口与拐点”", 13, G900, True)], anchor=MIDDLE)
    findings = (
        ("7 个子系统在 OOT 平台上直接拒绝（KV 分片、统一内存池、fp4 / mxfp8 KV pool 等）",
         "→ 插件或上游要补的接缝"),
        ("2 处要求插件实现 hook 才能用：CUDA graph runner、piecewise 编译后端", "→ 插件开发清单"),
        ("低 ridge point 芯片上，投机解码的 verify 在小 batch 就转为算力受限",
         "→ 给出收益消失的 batch 拐点"),
    )
    add_text(slide, rx + 0.22, ey + 0.6, rw - 0.44, eh - 0.7, [
        runs([(fact, 11, G700, False), (" " + result, 11, a, True)], bullet=a, after=7, spacing=1.1)
        for fact, result in findings
    ])
    footer(slide, 3, "虚线 = 规划中；数字来自内置能力快照（SGLang 72730d6，2026-10-14）")


# ---------------------------------------------------------------------- 4 loop B
def ellipse_point(cx, cy, rx, ry, degrees):
    angle = math.radians(degrees)
    return cx + rx * math.cos(angle), cy + ry * math.sin(angle)


def ellipse_arc(slide, cx, cy, rx, ry, start, end, color, width=1.5):
    """Clockwise arc between two parametric angles of the ellipse, with an arrowhead at the end."""

    def visual(degrees):
        angle = math.radians(degrees)
        return math.degrees(math.atan2(ry * math.sin(angle), rx * math.cos(angle))) % 360

    shape = slide.shapes.add_shape(
        MSO_SHAPE.ARC, Inches(cx - rx), Inches(cy - ry), Inches(2 * rx), Inches(2 * ry)
    )
    drop_theme_style(shape)
    # Preset arc angles are measured from the centre (not parametric), so a non-circular arc needs converting.
    shape.adjustments[0] = visual(start) * 0.6
    shape.adjustments[1] = visual(end) * 0.6
    shape.fill.background()
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    add_arrowheads(shape.line)
    return shape


def cycle_ring(slide, cx, cy, rx, ry, steps, hub, color, tint):
    """steps: (name, detail, planned) placed clockwise from the top, joined by arrowed arcs."""
    pill_w, pill_h = 1.9, 0.74
    hub_r = 0.86
    add_box(slide, cx - hub_r, cy - hub_r, 2 * hub_r, 2 * hub_r, fill=tint, kind=MSO_SHAPE.OVAL)
    add_text(slide, cx - hub_r, cy - hub_r, 2 * hub_r, 2 * hub_r,
             [P(hub[0], 11, color, True, align=CENTER, after=3)]
             + [P(line, 13, G900, True, align=CENTER) for line in hub[1:]], anchor=MIDDLE)
    step = 360 / len(steps)
    angles = [-90 + index * step for index in range(len(steps))]
    rects = []
    for index, ((name, detail, planned), angle) in enumerate(zip(steps, angles)):
        x, y = ellipse_point(cx, cy, rx, ry, angle)
        rects.append((x - pill_w / 2, y - pill_h / 2, x + pill_w / 2, y + pill_h / 2))
        pill = add_box(slide, x - pill_w / 2, y - pill_h / 2, pill_w, pill_h, fill=WHITE,
                       line=G400 if planned else color, line_w=1.25, dash=DASH if planned else None,
                       kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.35)
        boxed_text(pill, [
            runs([(f"{index + 1} ", 12.5, color, True, "Arial"), (name, 12.5, G900, True)],
                 align=CENTER, after=2),
            P(detail, 10, G500, align=CENTER),
        ], inset=(0.05, 0.02, 0.05, 0.02))

    def covered(px, py, rect):
        x0, y0, x1, y1 = rect
        return x0 - 0.09 <= px <= x1 + 0.09 and y0 - 0.09 <= py <= y1 + 0.09

    for index, angle in enumerate(angles):
        neighbours = (rects[index], rects[(index + 1) % len(rects)])
        free = [
            angle + step * sample / 200
            for sample in range(201)
            if not any(covered(*ellipse_point(cx, cy, rx, ry, angle + step * sample / 200), rect)
                       for rect in neighbours)
        ]
        if free:
            ellipse_arc(slide, cx, cy, rx, ry, free[0], free[-1], color, 1.5)


def mechanism(slide, x, y, w, number, title, body, color):
    number_dot(slide, x, y + 0.02, 0.34, color, number, 11)
    add_text(slide, x + 0.48, y, w - 0.48, 0.36, [P(title, 13, G900, True)], anchor=MIDDLE)
    add_text(slide, x + 0.48, y + 0.38, w - 0.48, 1.0, [P(body, 11, G700, spacing=1.12)])


def loop_b_slide(slide):
    b, tint = LOOPS["B"]["main"], LOOPS["B"]["tint"]
    header(slide, "LOOP B · 助力模型适配与评估优化",
           "新模型上 M100：从“能不能跑”一路推到“单机 SLO 内扛多少”",
           "关键步骤都产出可执行命令；实测数据回来后校准成本模型、重新规划", loop="B")
    cycle_ring(slide, 3.72, 4.42, 2.5, 1.9, [
        ("适配分析", "架构 · hook · 算子缺口", False),
        ("精度路线", "按精度算几卡放得下", False),
        ("特性组合", "接 Loop A 的 launch", False),
        ("仿真评估", "TTFT · TPOT · SLO", False),
        ("压测验证", "目前人工执行命令", True),
        ("校准回灌", "拟合系数 · 重新规划", False),
    ], ("模型级目标", "单机 SLO 内", "能扛多少请求"), b, tint)

    x, w = 7.25, SW - MX - 7.25
    mechanism(slide, x, 1.98, w, "1", "候选方案怎么来",
              "先套 Loop A 调好的 launch；TTFT 先超标时试 prefill 更快的布局（prefill CP、关 DP attention），"
              "TPOT 先超标且没开投机解码时试 mixed chunk（1024 / 512 / 256）", b)
    mechanism(slide, x, 3.27, w, "2", "怎么算",
              "仿真 SGLang 调度：prefill 优先、chunked prefill、KV 准入与回收、DP rank 同步走 step；"
              "按 P90 判 SLO，二分搜索单机在 SLO 内的最大到达率", b)
    mechanism(slide, x, 4.56, w, "3", "怎么信",
              "server 日志拟合 prefill / decode 系数（实测 / 模型取中位数）；"
              "压测检查预测偏差是否在 ±20% 内", b)
    iy = 5.72
    add_box(slide, x, iy, w, 1.03, fill=tint)
    add_text(slide, x + 0.18, iy + 0.08, w - 0.3, 0.95, [
        P("与 Loop C 的接口", 12.5, b, True, after=4),
        runs([("→ ", 11, b, True), ("算子缺口按占 step 比例分级（≥10% 高、≥3% 中）交给 Loop C", 11, G700)],
             after=3),
        runs([("← ", 11, LOOPS["C"]["main"], True), ("算子实测耗时回来，校准成本模型后重新规划", 11, G700)]),
    ])
    footer(slide, 4, "虚线 = 规划中：自动拉起服务与压测（--allow-exec）、自动 A/B 调参")


# ---------------------------------------------------------------------- 5 loop C
def loop_c_slide(slide):
    c, tint = LOOPS["C"]["main"], LOOPS["C"]["tint"]
    header(slide, "LOOP C · 助力算子绑定评估与优化",
           "算子先查“绑没绑对”，再查“快不快”，按模型占比排优先级",
           "绑定 = 模型的每个算子在 SGLang 里最终分发到哪条实现。生成了 kernel，不等于模型用上了它", loop="C")
    lw = 5.85
    add_text(slide, MX, 1.9, lw, 0.4, [P("① 绑定评估：模型用上的是哪条实现", 13.5, c, True)], anchor=MIDDLE)
    tiers = (
        ("模型实际用到的算子族", "norm · RoPE · 激活 · attention · KV 写入 · MoE · 线性层 · 通信"),
        ("SGLang 分发入口", "算子分发类 · 注意力后端 · 量化方法 · 平台 hook"),
    )
    y = 2.38
    for name, detail in tiers:
        box = add_box(slide, MX, y, lw, 0.64, fill=G100)
        boxed_text(box, [P(name, 12, G900, True, after=2), P(detail, 10.5, G700)],
                   inset=(0.16, 0.04, 0.1, 0.04))
        add_arrow(slide, MX + lw / 2, y + 0.66, MX + lw / 2, y + 0.86, G400, 1.5)
        y += 0.9
    add_box(slide, MX, y, lw, 1.02, fill=WHITE, line=c, line_w=1.0)
    add_text(slide, MX + 0.16, y + 0.06, lw - 0.3, 0.3,
             [runs([("M100 上实际走的路径", 12, G900, True), ("　缺口按占 step 比例排序", 10.5, G500, False)])],
             anchor=MIDDLE)
    statuses = (
        (OK, "厂商实现"), (OK, "SGLang 内置实现"), (WARN, "回落 torch 实现"),
        (WARN, "能跑但有 host 同步"), (BAD, "无可用实现"), (BAD, "当前配置会报错"),
    )
    chip_w, chip_h, chip_gap = 1.79, 0.27, 0.09
    for index, (color, label) in enumerate(statuses):
        row, col = divmod(index, 3)
        status_chip(slide, MX + 0.14 + col * (chip_w + chip_gap), y + 0.4 + row * (chip_h + 0.07),
                    chip_w, chip_h, color, label)
    cy = y + 1.2
    add_box(slide, MX, cy, lw, 6.75 - cy, fill=tint)
    add_text(slide, MX + 0.18, cy + 0.08, lw - 0.36, 6.75 - cy - 0.1, [
        P("绑定陷阱：模型能跑通，注册的 kernel 却没被调用", 12.5, c, True, after=4),
        P("register_oot_forward 只认精确类，子类不会继承注册", 11, G700, bullet=c, after=2),
        P("类自己重写 forward() 又不走分发时，注册的 kernel 根本不会被调用", 11, G700, bullet=c, after=2),
        P("Agent 从 SGLang 源码解析类继承链（MRO），直接标出这类路径", 11, G700, bullet=c),
    ])

    rx = MX + lw + 0.4
    rw = SW - MX - rx
    add_text(slide, rx, 1.9, rw, 0.4, [P("② 性能评估与优化：慢在哪、值不值得改", 13.5, c, True)], anchor=MIDDLE)
    rows = (
        ("按真实形状算期望", "从 step 标注重建 batch 形状，用同一成本模型算各类 kernel 的期望耗时", False),
        ("定位异常", "远低于屋顶线 · kernel 间空隙 · 未走图 · 耗时离群 · rank 不均衡", False),
        ("匹配优化手段", "融合 · 通算重叠 · 多流 · L2 权重预取 · MoE 调优配置，按芯片能力取舍", False),
        ("量化收益", "每条发现折算成 ms/step 与占比（≥10% 高、≥3% 中），交给 Loop B", False),
        ("A/B 复测", "规划中：自动执行；目前按报告里的命令人工复测", True),
    )
    pitch, d = 0.71, 0.36
    y0 = 2.4
    add_line(slide, rx + d / 2, y0 + d / 2, rx + d / 2, y0 + 3 * pitch + d / 2, G300, 1.25)
    add_line(slide, rx + d / 2, y0 + 3 * pitch + d / 2, rx + d / 2, y0 + 4 * pitch, G300, 1.25, dash=DASH)
    for index, (name, detail, planned) in enumerate(rows):
        y = y0 + index * pitch
        if planned:
            add_box(slide, rx, y, d, d, fill=WHITE, kind=MSO_SHAPE.OVAL)
        number_dot(slide, rx, y, d, G400 if planned else c, str(index + 1), 11, dashed=planned)
        add_text(slide, rx + 0.5, y - 0.03, rw - 0.5, 0.3, [P(name, 12.5, G500 if planned else G900, True)],
                 anchor=MIDDLE)
        add_text(slide, rx + 0.5, y + 0.28, rw - 0.5, 0.26, [P(detail, 10.5, G700)], anchor=MIDDLE)
    by = 6.03
    add_box(slide, rx, by, rw, 0.72, fill=G050, line=G400, dash=DASH)
    add_text(slide, rx + 0.18, by + 0.04, rw - 0.3, 0.64, [
        P("与平台算子 Agent 的衔接（规划中）", 11.5, G900, True, after=2),
        P("平台 kernel → 核对是否被分发 → 真实形状复测 → Loop B 换算模型级收益", 10.5, G700),
    ], anchor=MIDDLE)
    footer(slide, 5, "状态取值与判定规则来自 SGLang 源码快照（72730d6）；虚线 = 规划中")


# ---------------------------------------------------------------------- 6 differentiation
def differentiation_slide(slide):
    header(slide, "差异 · 算子级 VS 模型级",
           "与平台算子 Agent 互补：它把 kernel 做快，我们把模型跑好",
           "我们不止做算子级：以算子级结果为输入，一路评估到 step、调度和单机 SLO")
    b, c = LOOPS["B"]["main"], LOOPS["C"]["main"]
    cols = ((MX, 1.75), (2.45, 5.35), (7.95, SW - MX - 7.95))
    owner_x, owner_w = cols[2]
    hy = 1.95
    add_text(slide, cols[0][0] + 0.1, hy, cols[0][1], 0.5, [P("层级", 11.5, G500, True)], anchor=MIDDLE)
    add_text(slide, cols[1][0] + 0.1, hy, cols[1][1], 0.5, [P("这一层要回答的问题", 11.5, G500, True)],
             anchor=MIDDLE)
    add_text(slide, owner_x, hy, 1.5, 0.5, [P("谁来做 · 怎么做", 11.5, G500, True)], anchor=MIDDLE)
    legend_x = owner_x + 1.45
    for color, label, width in ((b, "我们 · 模型级", 1.32), (G700, "平台算子 Agent · 算子级", 1.98)):
        add_box(slide, legend_x, hy + 0.18, 0.14, 0.14, fill=color, kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
        add_text(slide, legend_x + 0.2, hy, width - 0.2, 0.5, [P(label, 10, G700, True)], anchor=MIDDLE)
        legend_x += width
    layers = (
        ("L5", "单机 SLO", "单机在 SLO 内能扛多少请求？先超标的是 TTFT 还是 TPOT？",
         "单机部署仿真：二分搜索 SLO 内最大到达率", LOOPS["B"]["tint"]),
        ("L4", "调度与特性", "chunked prefill、DP attention、投机解码、HiCache 怎么开？",
         "特性顾问 + 调度仿真", LOOPS["B"]["tint"]),
        ("L3", "Step", "真实 batch 形状下谁占大头？融合边界、通算重叠、多流、预取？",
         "trace 期望对实测；评估融合、通算重叠、多流、预取", LOOPS["C"]["tint"]),
        ("L2", "算子绑定", "模型的每个算子分发到哪条实现？缺口在哪、影响多大？",
         "静态解析分发路径；缺口按占 step 比例排序", LOOPS["C"]["tint"]),
        ("L1", "Kernel 实现", "单个 kernel 的 tiling、访存、指令怎么写最快？", None, G100),
    )
    y0, row_h, gap = 2.55, 0.58, 0.07
    for index, (code, name, question, ours, fill) in enumerate(layers):
        y = y0 + index * (row_h + gap)
        cell = add_box(slide, cols[0][0], y, cols[0][1], row_h, fill=fill)
        boxed_text(cell, [runs([(code + "  ", 10, G500, True, MONO), (name, 12.5, G900, True)])],
                   inset=(0.12, 0, 0.06, 0))
        q = add_box(slide, cols[1][0], y, cols[1][1], row_h, fill=G050)
        boxed_text(q, [P(question, 11, G700)], inset=(0.14, 0, 0.1, 0))
        if code == "L1":
            half_w = (owner_w - 0.1) / 2
            platform = add_box(slide, owner_x, y, half_w, row_h, fill=G700)
            boxed_text(platform, [P("生成与调优 kernel", 11, WHITE, True, align=CENTER)], inset=(0.06, 0, 0.06, 0))
            mine = add_box(slide, owner_x + half_w + 0.1, y, half_w, row_h, fill=WHITE, line=b, line_w=1.0,
                           dash=DASH)
            boxed_text(mine, [P("以平台 kernel 为输入", 10.5, b, True, align=CENTER)], inset=(0.06, 0, 0.06, 0))
        else:
            mine = add_box(slide, owner_x, y, owner_w, row_h, fill=b)
            boxed_text(mine, [P(ours, 11, WHITE, True, align=CENTER)], inset=(0.1, 0, 0.1, 0))
    sy, sh = 5.92, 0.86
    half = (CW - 0.25) / 2
    for index, (title, lines) in enumerate((
        ("评价口径不同", (("算子级：", "kernel 耗时、离屋顶线还有多远"),
                     ("模型级：", "单机 SLO 内容量、TTFT / TPOT、SLO 达成率"))),
        ("协作接口（规划中）", (("我们 → 平台：", "按模型级收益排序的算子需求（目标 shape、绑定位置）"),
                         ("平台 → 我们：", "kernel；接入后由我们做模型级复测与校准"))),
    )):
        x = MX + index * (half + 0.25)
        add_box(slide, x, sy, half, sh, fill=G100, line=G400 if index else None, dash=DASH if index else None)
        add_text(slide, x + 0.18, sy + 0.04, half - 0.3, sh - 0.08,
                 [P(title, 12, G900, True, after=3)]
                 + [runs([(lead, 10.5, c if index else G900, True), (text, 10.5, G700)], after=1)
                    for lead, text in lines], anchor=MIDDLE)
    footer(slide, 6, "虚线 = 规划中：与平台算子 Agent 的自动对接")


# ---------------------------------------------------------------------- 7 deep dive: the difficulty
def share_bar(slide, x, y, w, h, label, shares):
    add_text(slide, x - 0.52, y - 0.04, 0.46, h + 0.08, [P(label, 10, G500, True, MONO, align=RIGHT)],
             anchor=MIDDLE)
    left = x
    for fraction, color, text in shares:
        width = w * fraction
        seg = add_box(slide, left, y, width, h, fill=color)
        if text:
            boxed_text(seg, [P(text, 9.5, WHITE, True, align=CENTER)], inset=(0, 0, 0, 0))
        left += width


def difficulty_slide(slide):
    b, c = LOOPS["B"]["main"], LOOPS["C"]["main"]
    header(slide, "深挖 · 难点",
           "难点：算子同样快 30%，单机容量可能只涨 2%，也可能涨 40%",
           "推演：Qwen3-32B · 单机 8 卡、单个 TP8 实例 · 成本模型 + 离散事件仿真；未经实测校准，只用来说明方法")
    label_w, cell_w, gap = 1.95, 1.5, 0.06
    xs = [MX + label_w + gap + i * (cell_w + gap) for i in range(3)]
    hy = 1.98
    add_text(slide, MX + 0.16, hy, label_w - 0.16, 0.52,
             [P("单机 SLO 内容量", 11, G700, True), P("最大到达率", 9.5, G500)], anchor=MIDDLE)
    for x, head in zip(xs, ("基线", "attention\n耗时 −30%", "GEMM\n耗时 −30%")):
        lines = head.split("\n")
        add_text(slide, x, hy, cell_w, 0.52, [P(line, 11, G700, True, align=CENTER) for line in lines],
                 anchor=MIDDLE)
    rows = (
        ("chat", ("输入 1–4K", "P90 TTFT ≤ 2 s", "P90 TPOT ≤ 50 ms"),
         (("2.69", " req/s", ("TTFT 先超标",), "base"),
          ("+1.7%", "", ("2.74 req/s", "仍是 TTFT 先超标"), "same"),
          ("+40%", "", ("3.75 req/s", "仍是 TTFT 先超标"), "hero"))),
        ("长上下文 RAG", ("输入 25–32K", "P90 TTFT ≤ 8 s", "P90 TPOT ≤ 60 ms"),
         (("5.8", " req/min", ("TTFT 先超标",), "base"),
          ("+21%", "", ("7.0 req/min", "转为 TPOT 先超标"), "shift"),
          ("+31%", "", ("7.6 req/min", "转为 TPOT 先超标"), "shift"))),
    )
    y, row_h = 2.58, 1.4
    for name, details, cells in rows:
        label = add_box(slide, MX, y, label_w, row_h, fill=G100)
        boxed_text(label, [P(name, 14, G900, True, after=5)] + [P(d, 10, G500, after=1) for d in details],
                   inset=(0.16, 0.04, 0.08, 0.04))
        for x, (value, unit, notes, kind) in zip(xs, cells):
            fill = {"base": G050, "same": WHITE, "shift": WHITE, "hero": LOOPS["B"]["tint"]}[kind]
            color = {"base": G900, "same": G400, "shift": b, "hero": b}[kind]
            add_box(slide, x, y, cell_w, row_h, fill=fill, line=G200 if kind in ("same", "shift") else None)
            add_text(slide, x, y + 0.18, cell_w, 0.6,
                     [runs([(value, 26, color, True, "Arial"), (unit, 11, color, True)], align=CENTER)],
                     anchor=MIDDLE)
            note_color = G500 if kind == "same" else G700
            paragraphs = [P(note, 10, note_color, align=CENTER, after=1) for note in notes[:-1]]
            last = notes[-1]
            paragraphs.append(P(last, 10, c if kind == "shift" else note_color, kind == "shift", align=CENTER))
            add_text(slide, x + 0.05, y + 0.82, cell_w - 0.1, 0.5, paragraphs, anchor=TOP)
        y += row_h + 0.1
    add_text(slide, MX, y - 0.02, xs[-1] + cell_w - MX, 0.5, [
        P("芯片参数暂用内置 H20 预设（公开规格 + 默认效率假设）；M100 规格写入 YAML 即可重跑", 10, G500, after=1),
        P("基线 = 单个 TP8 实例 + EAGLE3（接受率假设 0.7）；launch 固定，只改算子耗时", 10, G500),
    ], anchor=MIDDLE)

    rx = 7.5
    rw = SW - MX - rx
    add_text(slide, rx, 1.98, rw, 0.4, [P("为什么不能直接换算", 14, G900, True)], anchor=MIDDLE)
    reasons_y = (2.48, 4.08, 5.14)
    number_dot(slide, rx, reasons_y[0], 0.32, c, "1", 11)
    add_text(slide, rx + 0.45, reasons_y[0] - 0.03, rw - 0.45, 0.36, [P("占比随形状变", 12.5, G900, True)],
             anchor=MIDDLE)
    add_text(slide, rx + 0.45, reasons_y[0] + 0.34, rw - 0.45, 0.26,
             [P("prefill 里 attention 的占比：4K 输入约 7%，32K 约 38%", 11, G700)], anchor=MIDDLE)
    bar_x, bar_w = rx + 0.97, rw - 0.97
    share_bar(slide, bar_x, reasons_y[0] + 0.7, bar_w, 0.24, "4K",
              ((0.82, SHARE_GEMM, "GEMM 82%"), (0.07, SHARE_ATTN, ""), (0.11, SHARE_OTHER, "")))
    share_bar(slide, bar_x, reasons_y[0] + 1.0, bar_w, 0.24, "32K",
              ((0.55, SHARE_GEMM, "GEMM 55%"), (0.38, SHARE_ATTN, "attention 38%"), (0.07, SHARE_OTHER, "")))
    legend_y = reasons_y[0] + 1.28
    legend_x = bar_x
    for color, text, width in ((SHARE_GEMM, "GEMM", 0.8), (SHARE_ATTN, "attention", 1.02),
                               (SHARE_OTHER, "通信等其他", 1.1)):
        add_box(slide, legend_x, legend_y + 0.055, 0.13, 0.13, fill=color)
        add_text(slide, legend_x + 0.18, legend_y, width - 0.18, 0.24, [P(text, 9.5, G500)], anchor=MIDDLE)
        legend_x += width
    add_text(slide, legend_x, legend_y, bar_x + bar_w - legend_x, 0.24,
             [P("TP8 估算", 9.5, G400, align=RIGHT)], anchor=MIDDLE)
    for index, (title, body, body_h) in enumerate((
        ("先超标的指标决定哪段有用",
         ("chat 先超标的是 TTFT，prefill 约八成时间在 GEMM：GEMM 快 30%，容量 +40%"
          "（TTFT 含排队，提速被放大）；attention 只占约 7%，仅 +1.7%"), 0.7),
        ("瓶颈会转移", "RAG 两种提速后，先超标的都从 TTFT 变成 TPOT：下一轮先优化谁，要按新瓶颈重新排", 0.55),
    ), start=2):
        top = reasons_y[index - 1]
        number_dot(slide, rx, top, 0.32, c, str(index), 11)
        add_text(slide, rx + 0.45, top - 0.03, rw - 0.45, 0.36, [P(title, 12.5, G900, True)], anchor=MIDDLE)
        add_text(slide, rx + 0.45, top + 0.34, rw - 0.45, body_h, [P(body, 11, G700, spacing=1.1)])

    ty = 6.08
    add_box(slide, MX, ty, CW, 0.7, fill=G900)
    add_text(slide, MX + 0.3, ty, CW - 0.6, 0.7, [runs([
        ("结论：算子优化值不值得做、先做哪个，要放到", 13, WHITE, True),
        ("“负载 + SLO + 部署”", 13, AMBER, True),
        ("里算——这正是模型级 Agent 的工作", 13, WHITE, True),
    ])], anchor=MIDDLE)
    footer(slide, 7, "复现：deploy --model qwen3-32b --hardware h20 --sizes 8 --no-pd 取基线 launch；"
                     "launch 固定，用脚本把算子耗时 ×0.7 后重测单机容量")


# ---------------------------------------------------------------------- 8 deep dive: the method
def method_slide(slide):
    b, c = LOOPS["B"]["main"], LOOPS["C"]["main"]
    header(slide, "深挖 · 解法",
           "解法：一条从 kernel 算到单机容量的换算链，实测校准后再做决策",
           "五层共用同一套模型和参数：算子、特性、布局任何一处变化，都能一路传到 SLO 和单机容量")
    add_text(slide, MX, 1.88, 1.0, 0.3, [P("算子级", 10.5, c, True)], anchor=MIDDLE)
    add_text(slide, SW - MX - 1.0, 1.88, 1.0, 0.3, [P("模型级", 10.5, b, True, align=RIGHT)], anchor=MIDDLE)
    add_arrow(slide, MX + 0.75, 2.03, SW - MX - 0.75, 2.03, G300, 1.25)
    levels = (
        ("KERNEL", "算子耗时",
         ("roofline：", "max(算力, 访存)", "+ launch / 图间隙", "+ α-β 通信", "效率系数取自芯片 YAML"),
         c, LOOPS["C"]["tint"]),
        ("STEP", "单步耗时",
         ("按真实 batch 形状", "逐层累加各类算子", "prefill / decode / mixed", "trace 的形状从 step", "标注里重建"),
         c, LOOPS["C"]["tint"]),
        ("SCHEDULE", "请求时延",
         ("离散事件仿真 SGLang", "prefill 优先", "chunked prefill", "KV 准入与回收", "DP rank 同步走 step"),
         G700, G100),
        ("SLO", "SLO 判定",
         ("按 P90 判 TTFT / TPOT", "统计 SLO 达成率", "标出先超标的指标"),
         b, LOOPS["B"]["tint"]),
        ("CAPACITY", "单机容量",
         ("二分搜索 SLO 内", "最大到达率", "候选配置逐个比较", "取容量最高的"),
         b, LOOPS["B"]["tint"]),
    )
    bw, gap, by, bh = 2.18, 0.3, 2.3, 1.9
    centers = []
    for index, (code, title, lines, color, fill) in enumerate(levels):
        x = MX + index * (bw + gap)
        centers.append(x + bw / 2)
        add_box(slide, x, by, bw, bh, fill=fill)
        add_text(slide, x + 0.16, by + 0.14, bw - 0.24, bh - 0.2, [
            P(code, 10, color, True, MONO, after=3),
            P(title, 15, G900, True, after=7),
        ] + [P(line, 10.5, G700, after=2.5) for line in lines])
        if index:
            add_arrow(slide, x - gap + 0.05, by + bh / 2, x - 0.05, by + bh / 2, G500, 1.5)
    ly = by + bh + 0.3
    add_path_arrow(slide, [(centers[4], by + bh), (centers[4], ly), (centers[0], ly), (centers[0], by + bh)],
                   G500, 1.25)
    add_text(slide, centers[0] + 0.25, ly + 0.05, centers[4] - centers[0] - 0.5, 0.32, [runs([
        ("实测校准　", 10.5, G900, True),
        ("server 日志拟合 prefill / decode 系数；压测检查预测偏差是否在 ±20% 内；校准后重新规划",
         10.5, G700),
    ], align=CENTER)], anchor=MIDDLE)

    add_text(slide, MX, 5.1, CW, 0.36, [P("Agent 怎么用这条链", 13.5, G900, True)], anchor=MIDDLE)
    uses = (
        ("排优先级", "已实现", False, ("算子缺口、融合 / 重叠 / 异常等发现", "都折算成 ms/step 与占 step 比例来分级")),
        ("what-if", "规划中：做成命令", True, ("算子提速 → 重新仿真 → 单机容量变化", "上一页的推演就是用这条链手工跑的")),
        ("向 Loop C 与平台提需求", "规划中", True, ("从“把 X 做快”", "变成“这些 shape 下快多少、单机容量涨多少”")),
    )
    uw, ugap, uy, uh = (CW - 2 * 0.25) / 3, 0.25, 5.54, 1.12
    for index, (title, status, planned, lines) in enumerate(uses):
        x = MX + index * (uw + ugap)
        tag_w = 1.45 if planned and len(status) > 3 else 0.72
        add_box(slide, x, uy, uw, uh, fill=G100)
        add_text(slide, x + 0.16, uy + 0.12, uw - tag_w - 0.42, 0.3, [P(title, 12.5, G900, True)], anchor=MIDDLE)
        tag(slide, x + uw - 0.16 - tag_w, uy + 0.15, tag_w, 0.25, status, planned)
        add_text(slide, x + 0.16, uy + 0.5, uw - 0.26, uh - 0.56,
                 [P(line, 10.5, G700, after=3) for line in lines])
    footer(slide, 8)


# ---------------------------------------------------------------------- 9 closing
def status_list(slide, x, y, w, label, items, color, text_color, planned):
    add_text(slide, x, y, w, 0.3, [P(label, 11, color, True)], anchor=MIDDLE)
    d = 0.12
    for index, item in enumerate(items):
        top = y + 0.38 + index * 0.37
        add_box(slide, x + 0.02, top + (0.3 - d) / 2, d, d, fill=None if planned else color,
                line=color if planned else None, line_w=1.25, kind=MSO_SHAPE.OVAL)
        add_text(slide, x + 0.28, top, w - 0.28, 0.3, [P(item, 11.5, text_color)], anchor=MIDDLE)


def closing_slide(slide):
    set_background(slide, G900)
    add_text(slide, MX, 0.45, 8.0, 0.32, [P("现状与下一步", 11, AMBER, True, MONO)], anchor=MIDDLE)
    add_text(slide, MX, 0.82, CW, 0.62, [P("评估与决策已打通，下一步把执行与复测自动化", 27, WHITE, True)],
             anchor=MIDDLE)
    columns = (
        ("A", "Loop A  框架迭代",
         ("源码能力快照（sync-knowledge）", "特性顾问：判可行、算收益、给参数", "插件 hook 清单与合并 launch"),
         ("两版快照自动 diff", "插件脚手架生成", "特性 A/B 自动执行")),
        ("B", "Loop B  模型适配与评估优化",
         ("适配分析 · 精度路线 · bring-up 计划", "单机部署仿真 · launch / 压测命令", "日志 / 压测校准 · 一键 run"),
         ("自动拉起服务与压测（--allow-exec）", "自动 A/B 调参")),
        ("C", "Loop C  算子绑定评估与优化",
         ("算子绑定解析与缺口排序", "roofline 与 trace 诊断", "融合 / 通算重叠 / 多流 / L2 预取评估"),
         ("优化项 A/B 自动复测", "what-if：算子提速 → 单机容量", "对接平台算子 Agent")),
    )
    cw, gap, y, h = (CW - 2 * 0.28) / 3, 0.28, 1.85, 4.2
    for index, (letter, title, done, planned) in enumerate(columns):
        x = MX + index * (cw + gap)
        color = LOOPS[letter]["dark"]
        add_box(slide, x, y, cw, h, fill=G800)
        ring_badge(slide, x + 0.42, y + 0.42, 0.5, color, letter, 13, width=2.0)
        add_text(slide, x + 0.8, y + 0.2, cw - 0.9, 0.44, [P(title, 13.5, WHITE, True)], anchor=MIDDLE)
        status_list(slide, x + 0.24, y + 0.88, cw - 0.4, "已实现", done, LOOPS["B"]["dark"], G200, False)
        add_line(slide, x + 0.24, y + 2.42, x + cw - 0.24, y + 2.42, "3A4755", 0.75)
        status_list(slide, x + 0.24, y + 2.58, cw - 0.4, "规划中", planned, AMBER, G300, True)
    add_text(slide, MX, 6.3, CW, 0.5, [runs([
        ("定位：不替代算子生成，而是把算子能力兑现为", 17, WHITE, True),
        ("模型级收益", 17, AMBER, True),
    ])], anchor=MIDDLE)
    footer(slide, 9, dark=True)


def add_notes(slide, text):
    frame = slide.notes_slide.notes_text_frame
    frame.text = text
    for index, paragraph in enumerate(frame.paragraphs):
        paragraph.space_after = Pt(6)
        for run in paragraph.runs:
            style_run(run, 12, G500 if index == 0 else G900, bold=index == 0)


def build(path):
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    layout = prs.slide_layouts[6]
    makers = (cover_slide, overview_slide, loop_a_slide, loop_b_slide, loop_c_slide,
              differentiation_slide, difficulty_slide, method_slide, closing_slide)
    for page, make in enumerate(makers, start=1):
        slide = prs.slides.add_slide(layout)
        make(slide)
        add_notes(slide, notes_text(page))
    prs.save(path)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "sglang_maestro_agent_loops.pptx")
