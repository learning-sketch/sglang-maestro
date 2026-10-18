"""Flag text that may wrap badly or overflow in PowerPoint/Keynote, whose CJK fonts set text wider than LibreOffice.

Widths are estimated conservatively: CJK characters and full-width punctuation take 1 em, Latin
glyphs use PingFang/YaHei-like advances. A paragraph is flagged when it wraps and its last line is
nearly empty (an orphan), or when a frame's estimated text height exceeds the frame.
"""

import math
import sys
import unicodedata

from pptx import Presentation

EMU_PER_PT = 12700
SAFETY = 0.95
LINE = 1.22

NARROW = {" ": 0.3, ".": 0.3, ",": 0.3, ":": 0.3, ";": 0.3, "'": 0.25, "|": 0.3, "!": 0.3,
          "(": 0.38, ")": 0.38, "[": 0.38, "]": 0.38, "/": 0.4, "-": 0.4, "_": 0.55,
          "+": 0.62, "=": 0.62, "<": 0.62, ">": 0.62, "%": 0.88, "&": 0.72, "#": 0.62,
          "*": 0.5, "~": 0.62, "^": 0.6, "×": 0.62, "·": 0.6, "±": 0.62, "≤": 0.8, "≥": 0.8,
          "α": 0.6, "β": 0.6, "−": 0.62, "–": 0.55}


def char_em(ch, bold):
    if ch in NARROW:
        return NARROW[ch]
    if ch.isascii():
        if ch.isdigit():
            em = 0.6
        elif ch.isupper():
            em = 0.7
        elif ch in "ijlrtf":
            em = 0.36
        elif ch in "mw":
            em = 0.88
        else:
            em = 0.58
        return em * (1.06 if bold else 1.0)
    if unicodedata.east_asian_width(ch) in ("W", "F", "A") or ch in "“”‘’→←↑↓…—○●✓":
        return 1.0
    return 0.65


def paragraph_runs(paragraph):
    for run in paragraph.runs:
        size = run.font.size.pt if run.font.size is not None else 18.0
        bold = bool(run.font.bold)
        yield run.text, size, bold


def check(path, only=None):
    prs = Presentation(path)
    issues = 0
    for number, slide in enumerate(prs.slides, start=1):
        if only and number not in only:
            continue
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            frame = shape.text_frame
            body = frame._txBody.find("{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr")
            insets = [body.get(key) for key in ("lIns", "tIns", "rIns", "bIns")]
            defaults = (91440, 45720, 91440, 45720)
            left, top, right, bottom = (int(v) if v is not None else d for v, d in zip(insets, defaults))
            avail_w = (shape.width - left - right) / EMU_PER_PT * SAFETY
            avail_h = (shape.height - top - bottom) / EMU_PER_PT
            total_h = 0.0
            for paragraph in frame.paragraphs:
                parts = list(paragraph_runs(paragraph))
                if not parts:
                    continue
                ppr = paragraph._p.pPr
                indent = int(ppr.get("marL", "0")) / EMU_PER_PT if ppr is not None else 0.0
                width = sum(char_em(ch, bold) * size for text, size, bold in parts for ch in text)
                size = max(size for _, size, _ in parts)
                spacing = paragraph.line_spacing if isinstance(paragraph.line_spacing, float) else 1.0
                usable = max(avail_w - indent, 1.0)
                lines = max(1, math.ceil(width / usable))
                before = paragraph.space_before.pt if paragraph.space_before is not None else 0.0
                after = paragraph.space_after.pt if paragraph.space_after is not None else 0.0
                total_h += lines * size * LINE * spacing + before + after
                text = "".join(t for t, _, _ in parts)
                if lines > 1:
                    tail = width - (lines - 1) * usable
                    if tail < max(2.6 * size, 0.12 * usable):
                        issues += 1
                        print(f"slide {number}: ORPHAN ({lines} lines, tail {tail / size:.1f} em) "
                              f"w={usable / 72:.2f}in :: {text}")
            if total_h > avail_h * 1.04 + 2:
                issues += 1
                text = frame.text.replace("\n", " / ")
                print(f"slide {number}: OVERFLOW est {total_h / 72:.2f}in > box {avail_h / 72:.2f}in :: {text[:90]}")
    print(f"{issues} potential issue(s)")


if __name__ == "__main__":
    only = {int(item) for item in sys.argv[2].split(",")} if len(sys.argv) > 2 else None
    check(sys.argv[1], only)
