"""Draft of the iteration page: who writes the patch, and the ladder one kernel chain went down.

The method page ends at "the ledger ranked four targets"; this page answers what happens next.
The top strip is the loop; under it the hard role split -- the subagent proposes and writes,
the controller builds and judges, and the mailbox between them carries an identity the reply
must match. The ladder is the payoff: 148.7 to 16.8 microseconds in five accepted steps, each
one re-measured on its own isolated chain, with the end-to-end re-verification beside it.

    python draft_iterate_page.py [out.pptx]
"""

import sys

from build_agent_deck import (
    CW,
    DASH,
    FONT,
    MIDDLE,
    MONO,
    MX,
    RIGHT,
    SH,
    SW,
    WHITE,
    P,
    add_arrow,
    add_box,
    add_line,
    add_text,
    runs,
)
from build_model_layer_page import THEMES
from pptx import Presentation
from pptx.util import Inches

T = THEMES["review"]
TITLE_BAR, TITLE_INK = "306FFF", "000000"
OUTPUT = "../../slides/iterate_page_draft.pptx"

TITLE = "M100 Agent：找到之后，改是怎么发生的"
LOOP = ("trace → 归因 → 账本排序 → 提假设 → 证据检查 → 隔离补丁 + 登记构建 → C 筛选 → "
        "数值与质量门 → D 三方复验 → 晋升 → 回到 trace")

# (title, subtitle, accent colour key, (line, ...), dashed)
ROLES = (
    ("Codex subagent", "/root/optimization_planner", "accent", (
        "读真实源码与冻结记录，形成假设",
        "提出结构化动作：读源码 / 补丁 / 微实验 / 候选",
        "本轮唯一规划器，补丁由它自己写",
    ), False),
    ("planner_mailbox", "请求与回复都带身份", "g400", (
        "单次请求含 UUID + 请求 SHA + 动作 schema",
        "回复须与请求 ID / SHA 匹配并提供来源",
        "请求、回复、失败全部落盘",
    ), True),
    ("Python 控制器", "落地与判决", "ship", (
        "设备锁、状态机、源码工作区、构建、服务启停",
        "三道门、错误记录、best 晋升",
        "候选的成败由它说了算",
    ), False),
)
CONSTRAINT = "规划器不能自行宣布实验成功：数值、图文质量与性能收益，全部由控制器独立检查"

BASELINE = ("基线", "", "原始 Q/K Norm + MRoPE 路径", 148.687)
# (step, candidates, what changed, before us, after us)
LADDER = (
    ("1", "001→002", "Torch FP32 统计 + native MRoPE", 148.687, 120.969),
    ("2", "003→004", "native tail", 120.585, 65.820),
    ("3", "009→010", "native stats + Torch rsqrt，支持真实 stride", 65.816, 30.221),
    ("4", "012→013", "单核 mode3", 30.226, 20.067),
    ("5", "025→026", "group4：4 head/CTA，warp 不变、减发射数", 20.047, 16.839),
)
SIDELINE = ("旁线", "020→021", "开 klx_moe_prepare_input 插件", 240.338, 216.896)

VERIFY = (
    ("008", "002+004", (
        "c128 1399.47 → 1504.50 tok/s（+7.51%）",
        "c32 名义 +3.02%，噪声门 4.84%，不单独确认",
    )),
    ("019", "010+013", (
        "全 5 档同库对照，三道门全过",
        "c1 +50.0% · c32 +37.1% · c128 +35.9%",
    )),
    ("034", "021+026", (
        "按授权只测 c32：1409.66 tok/s",
        "相对原始 +45.61%，相对上一验收版 +6.96%",
    )),
)
BAND = "RoPE 链 7.196 ms / 912 个 kernel → 0.723 ms / 48 个；真实首层 19 个 kernel 收成 1 个"
FOOT = (
    "native 构建只编译登记过的 standalone 算子组件，未重编整个 xInfer 库；SGLang 与 xInfer "
    "用原仓库原 Git 版本，不另起开发树",
    "这个 native kernel 由 subagent 自己写，不是 xpugen 生成；规划器依赖当前 Codex 会话托管，"
    "独立跑控制器还不能自行拉起，长跑能力尚未闭环",
)
ROW_H = 0.26


def tone(key):
    return {"accent": T.accent, "ship": T.ship, "g400": T.g400}[key]


def header(slide):
    add_box(slide, MX, 0.44, 0.07, 0.38, fill=TITLE_BAR)
    add_text(slide, MX + 0.22, 0.33, 10.4, 0.6, [P(TITLE, 26, TITLE_INK, True)], anchor=MIDDLE)


def loop_strip(slide, y):
    add_box(slide, MX, y, CW, 0.34, fill=T.panel)
    add_text(slide, MX + 0.2, y, CW - 0.4, 0.34, [runs([
        ("一轮闭环　", 10.5, T.g900, True), (LOOP, 9.5, T.g700),
    ])], anchor=MIDDLE)


def role_box(slide, x, w, y, h, role):
    title, subtitle, key, lines, dashed = role
    color = tone(key)
    add_box(slide, x, y, w, h, fill=T.panel,
            line=color if dashed else None, dash=DASH if dashed else None)
    if not dashed:
        add_box(slide, x, y, 0.06, h, fill=color)
    add_text(slide, x + 0.24, y + 0.06, w - 0.44, 0.24, [runs([
        (title, 11.5, color, True, MONO if key == "g400" else FONT),
    ])], anchor=MIDDLE)
    add_text(slide, x + 0.24, y + 0.28, w - 0.44, 0.20,
             [P(subtitle, 9, T.g400, False, MONO)], anchor=MIDDLE)
    for index, line in enumerate(lines):
        add_text(slide, x + 0.24, y + 0.50 + index * 0.20, w - 0.44, 0.20,
                 [P(line, 9.5, T.g700)], anchor=MIDDLE)


def ladder_row(slide, x, w, top, scale, step, cands, text, before, after, muted=False):
    bar_x, bar_w = x + 4.32, 1.21
    ink = T.g500 if muted else T.g900
    add_text(slide, x + 0.22, top, 1.05, ROW_H, [runs([
        (step, 10, T.g500 if muted else T.accent, True),
        (f"　{cands}", 8.5, T.g400, False, MONO),
    ])], anchor=MIDDLE)
    add_text(slide, x + 1.32, top, 2.95, ROW_H, [P(text, 9.5, ink)], anchor=MIDDLE)
    if scale:
        bar_h = 0.13
        add_box(slide, bar_x, top + (ROW_H - bar_h) / 2, max(0.02, after * scale), bar_h,
                fill=T.accent)
        if before * scale > after * scale + 0.04:
            add_box(slide, bar_x + after * scale, top + (ROW_H - bar_h) / 2,
                    before * scale - after * scale, bar_h, line=T.g400, line_w=0.75, dash=DASH)
    else:
        add_text(slide, bar_x, top, bar_w, ROW_H,
                 [P("另一条链", 8.5, T.g400)], anchor=MIDDLE)
    span = f"{after:.1f}" if before == after else f"{before:.1f} → {after:.1f}"
    gain = "—" if before == after else f"+{before / after - 1:.0%}"
    add_text(slide, x + 5.61, top, 1.02, ROW_H,
             [P(span, 9, ink, False, MONO, align=RIGHT)], anchor=MIDDLE)
    add_text(slide, x + 6.68, top, 0.60, ROW_H,
             [P(gain, 9.5, T.g400 if muted else ink, not muted, MONO, align=RIGHT)],
             anchor=MIDDLE)


def ladder_panel(slide, x, y, w, h):
    add_box(slide, x, y, w, h, fill=T.panel)
    add_text(slide, x + 0.22, y + 0.06, w - 0.44, 0.28, [runs([
        ("一个算子链的五步", 12.5, T.g900, True),
        (f"　148.7 → 16.8 µs，合计 {148.687 / 16.839:.1f} 倍 · 每步都在隔离链上精确重测",
         10, T.g500),
    ])], anchor=MIDDLE)
    heads = ((x + 0.22, 1.05, None, "步 · 候选"), (x + 1.32, 2.95, None, "做了什么"),
             (x + 4.32, 1.21, None, "耗时 · 虚线=省掉"), (x + 5.61, 1.02, RIGHT, "µs"),
             (x + 6.68, 0.60, RIGHT, "加速"))
    for cell_x, cell_w, align, head in heads:
        add_text(slide, cell_x, y + 0.38, cell_w, 0.22,
                 [P(head, 9, T.g400, align=align) if align else P(head, 9, T.g400)],
                 anchor=MIDDLE)
    add_line(slide, x + 0.22, y + 0.62, x + w - 0.22, y + 0.62, T.g200, 0.75)

    scale = 1.21 / BASELINE[3]
    top = y + 0.66
    label, cands, text, value = BASELINE
    ladder_row(slide, x, w, top, scale, label, cands, text, value, value, muted=True)
    top += ROW_H
    for step, cands, text, before, after in LADDER:
        ladder_row(slide, x, w, top, scale, step, cands, text, before, after)
        top += ROW_H
    add_line(slide, x + 0.22, top + 0.03, x + w - 0.22, top + 0.03, T.g200, 0.75)
    ladder_row(slide, x, w, top + 0.06, None, *SIDELINE, muted=True)


def verify_panel(slide, x, y, w, h):
    add_box(slide, x, y, w, h, fill=T.panel)
    add_text(slide, x + 0.22, y + 0.06, w - 0.44, 0.28, [runs([
        ("三轮端到端复验", 12.5, T.g900, True), ("　隔离链的收益不算数", 10, T.g500),
    ])], anchor=MIDDLE)
    top = y + 0.40
    for package, combo, lines in VERIFY:
        add_box(slide, x + 0.16, top, w - 0.32, 0.66, fill=WHITE, line=T.g200)
        add_text(slide, x + 0.32, top + 0.04, w - 0.64, 0.22, [runs([
            (package, 11, T.accent, True, MONO), (f"　{combo}", 9.5, T.g500, False, MONO),
        ])], anchor=MIDDLE)
        for index, line in enumerate(lines):
            add_text(slide, x + 0.32, top + 0.26 + index * 0.19, w - 0.64, 0.19,
                     [P(line, 9, T.g700)], anchor=MIDDLE)
        top += 0.72


def iterate_page(slide):
    header(slide)
    loop_strip(slide, 1.04)

    roles_y, roles_h, gap = 1.52, 1.12, 0.26
    role_w = (CW - gap * (len(ROLES) - 1)) / len(ROLES)
    for index, role in enumerate(ROLES):
        x = MX + index * (role_w + gap)
        role_box(slide, x, role_w, roles_y, roles_h, role)
        if index:
            add_arrow(slide, x - gap + 0.04, roles_y + roles_h / 2, x - 0.04,
                      roles_y + roles_h / 2, T.g400, 1.25)
    add_text(slide, MX, roles_y + roles_h + 0.06, CW, 0.24, [runs([
        ("硬约束　", 10, T.drop, True), (CONSTRAINT, 10, T.g700),
    ])], anchor=MIDDLE)

    panel_y, panel_h, ladder_w = 3.00, 2.64, 7.5
    ladder_panel(slide, MX, panel_y, ladder_w, panel_h)
    verify_x = MX + ladder_w + 0.2
    verify_panel(slide, verify_x, panel_y, SW - MX - verify_x, panel_h)

    T.band(slide, 5.78, 0.54, "读数", BAND)
    add_text(slide, MX, 6.44, 0.5, 0.24, [P("边界", 10, T.accent, True)], anchor=MIDDLE)
    add_text(slide, MX + 0.5, 6.44, CW - 0.5, 0.48,
             [P(line, 9.5, T.g500) for line in FOOT])


def build(path):
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    iterate_page(prs.slides.add_slide(prs.slide_layouts[6]))
    prs.save(path)


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else OUTPUT
    build(target)
    print(target)
