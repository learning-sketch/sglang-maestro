"""Draft of the overall-plan page: two agents, one skeleton, drawn so the rest of the deck nests
inside it.

The page used to put the two layers side by side with equal detail -- gates, measurement scope,
iteration budget -- which made it a comparison table rather than an overview, and left the claim
that the two layers are the same shape to be inferred from symmetry. Here the claim comes first
and the two layers are drawn as the same four-part run, so a reader sees one method twice
instead of two methods once.

Stacked rather than columned: a 16:9 page gives each row the full width its boxes need, and the
two couplings between the layers (a bottleneck operator going down, a new kernel coming back up)
then run vertically, which is the direction they actually go.

Colour carries one distinction and nothing else: orange where a model is in the loop, slate
where the verdict is rules only. The note under a box names what expands it later, so the
hand-off to the following pages is made by recognition rather than by printed page numbers.

    python draft_overview_page.py [out.pptx]
"""

import sys

from build_agent_deck import (
    CENTER,
    CW,
    FONT,
    G050,
    G200,
    G500,
    G700,
    G900,
    LOOPS,
    MIDDLE,
    MONO,
    MX,
    RIGHT,
    SH,
    SW,
    WHITE,
    P,
    add_arrow,
    add_box,
    add_path_arrow,
    add_text,
    runs,
)
from pptx import Presentation
from pptx.util import Inches

AGENT = LOOPS["C"]["main"]
TINT = LOOPS["C"]["tint"]
RULES = G700
OUTPUT = "../../slides/overview_page_draft.pptx"
TITLE = "M100 Agent：总体方案"
SUBTITLE = "两个 Agent，同一套骨架：模型负责生成方案，判决交给不含模型的固定规则与真机实测"

ROW_H = 2.18
ROW_Y = (1.24, 4.30)
EX_Y, EX_H = 3.54, 0.64
KEY_Y, KEY_H = 6.62, 0.58

# Four parts of one run, at the same x in both rows: who asks, who generates, who judges, who
# takes delivery. (x, w) -- the two middle boxes are the ones the later pages open up.
SLOTS = ((0.80, 1.55), (2.75, 3.45), (6.60, 4.05), (11.05, 1.50))

# (eyebrow, layer, role, ask, generate, judge, take, return label, cases)
# generate / judge are (title, lines, note); note is what expands it later, or None.
ROWS = (
    ("model-agent", "模型层", "已接入模型在真实服务下的端到端吞吐",
     ("人", "定目标与边界"),
     ("Agent 生成方案", ("提假设、改代码",), "人与 Agent 的分工"),
     ("规则验收 + 真机实测", ("整模型端到端", "真实服务 + 固定评测集"), "八步运行机制与三道门"),
     ("人", "判断交付"),
     "不过即淘汰，带着原因退回",
     "Qwen3-VL-30B-A3B"),
    ("xpugen", "算子层", "面向 M100 自动生成算子实现并调优",
     ("需求", "输入规格与约束"),
     ("Agent 写算子", ("全流程唯一的模型",), None),
     ("Harness 判算子", ("规则驱动、无模型", "单算子、多 Shape，与 Golden 对齐"),
      "L0 编译运行 · L1 与 Golden 对齐 · Bench 真机计时"),
     ("候选", "交付实现"),
     "不过即淘汰，带着证据退回",
     "batch_norm"),
)

EXCHANGE = (
    (1.05, False, "模型层 → 算子层", "Trace 定位的瓶颈算子，转成带输入规格的算子请求"),
    (6.95, True, "算子层 → 模型层", "新内核合入候选分支，重新走一遍模型层的判决"),
)
LEGEND = ((AGENT, "有模型参与"), (RULES, "纯规则判决"))
KEY = ("两层唯一的区别，是判决在什么粒度上做",
       "算子 bench 只证明局部更快；能不能交付，由整模型复验决定")


def header(slide):
    add_box(slide, MX, 0.30, 0.07, 0.38, fill=AGENT)
    add_text(slide, MX + 0.24, 0.19, 8.0, 0.6, [P(TITLE, 25, G900, True)], anchor=MIDDLE)
    add_text(slide, MX + 0.02, 0.80, 8.6, 0.3, [P(SUBTITLE, 12, G700)], anchor=MIDDLE)
    # Colour says one thing on this page, so it is worth naming once.
    x = SW - MX
    for tone, label in reversed(LEGEND):
        label_w = len(label) * 10 / 72 + 0.1
        x -= 0.18 + label_w
        add_box(slide, x, 0.88, 0.12, 0.12, fill=tone)
        add_text(slide, x + 0.2, 0.80, label_w, 0.28, [P(label, 10, G500)], anchor=MIDDLE)
        x -= 0.22


def cap(slide, x, w, y, h, name, detail):
    """An end of the run: who hands work in, who takes delivery. Outside the automated part."""
    add_box(slide, x, y, w, h, fill=G050, line=G200)
    add_text(slide, x + 0.14, y + 0.04, w - 0.28, h - 0.08, [
        P(name, 12.5, G900, True, align=CENTER),
        P(detail, 9.5, G500, align=CENTER),
    ], anchor=MIDDLE)


def step(slide, x, w, y, h, tone, title, lines):
    add_box(slide, x, y, w, h, fill=WHITE, line=G200)
    add_box(slide, x, y, 0.06, h, fill=tone)
    add_text(slide, x + 0.2, y + 0.04, w - 0.32, h - 0.08, [
        P(title, 12.5, G900, True, after=2),
        *[P(line, 10, G700) for line in lines],
    ], anchor=MIDDLE)


def row(slide, index):
    (eyebrow, layer, role, ask, generate, judge, take, back, case) = ROWS[index]
    y = ROW_Y[index]
    add_box(slide, MX, y, CW, ROW_H, fill=WHITE, line=G200)
    add_box(slide, MX, y, CW, 0.05, fill=AGENT)
    add_text(slide, MX + 0.2, y + 0.08, CW - 0.4, 0.3, [runs([
        (eyebrow, 11.5, AGENT, True, MONO), ("　｜　", 11.5, G200),
        (layer, 13, G900, True), (f"　{role}", 10.5, G500),
    ])], anchor=MIDDLE)

    box_y, box_h = y + 0.74, 0.92
    cap(slide, SLOTS[0][0], SLOTS[0][1], box_y, box_h, *ask)
    step(slide, SLOTS[1][0], SLOTS[1][1], box_y, box_h, AGENT, generate[0], generate[1])
    step(slide, SLOTS[2][0], SLOTS[2][1], box_y, box_h, RULES, judge[0], judge[1])
    cap(slide, SLOTS[3][0], SLOTS[3][1], box_y, box_h, *take)
    for left, right in zip(SLOTS, SLOTS[1:]):
        add_arrow(slide, left[0] + left[1], box_y + box_h / 2, right[0],
                  box_y + box_h / 2, G500, 1.5)

    # The one line that makes this a loop and not a pipeline: a rejected candidate goes back to
    # the party that proposed it, carrying the reason with it.
    gen_mid = SLOTS[1][0] + SLOTS[1][1] / 2
    judge_mid = SLOTS[2][0] + SLOTS[2][1] / 2
    back_y = y + 0.54
    add_path_arrow(slide, [(judge_mid, box_y), (judge_mid, back_y), (gen_mid, back_y),
                           (gen_mid, box_y)], AGENT, 1.25)
    label_w = 3.0
    add_box(slide, (gen_mid + judge_mid) / 2 - label_w / 2, back_y - 0.14, label_w, 0.28,
            fill=WHITE)
    add_text(slide, (gen_mid + judge_mid) / 2 - label_w / 2, back_y - 0.14, label_w, 0.28,
             [P(back, 9.5, AGENT, align=CENTER)], anchor=MIDDLE)

    note_y = y + 1.72
    for slot, part in ((SLOTS[1], generate), (SLOTS[2], judge)):
        if part[2]:
            add_text(slide, slot[0] + 0.06, note_y, 5.0, 0.3, [runs([
                ("后面展开　", 9, AGENT, True), (part[2], 9.5, G700),
            ])], anchor=MIDDLE)
    add_text(slide, SLOTS[3][0] - 2.0, note_y, SLOTS[3][1] + 2.0, 0.3, [runs([
        ("案例　", 9, G500), (case, 9.5, G700, False, MONO),
    ], align=RIGHT)], anchor=MIDDLE)


def exchange(slide):
    """The only two couplings between the layers -- and the only reason this is one system."""
    add_box(slide, MX, EX_Y, CW, EX_H, fill=TINT)
    for x, upward, lead, text in EXCHANGE:
        top, bottom = EX_Y + 0.14, EX_Y + EX_H - 0.14
        add_arrow(slide, x, bottom if upward else top, x, top if upward else bottom, AGENT, 1.5)
        add_text(slide, x + 0.22, EX_Y, 5.3, EX_H, [runs([
            (f"{lead}　", 10, AGENT, True), (text, 10.5, G700),
        ])], anchor=MIDDLE)


def key_band(slide):
    add_box(slide, MX, KEY_Y, CW, KEY_H, fill=TINT)
    add_box(slide, MX, KEY_Y, 0.07, KEY_H, fill=AGENT)
    statement, note = KEY
    add_text(slide, MX + 0.34, KEY_Y + 0.04, CW - 0.7, KEY_H - 0.08, [
        runs([("关键　", 11.5, AGENT, True, MONO), (statement, 14, G900, True)]),
        P(note, 11, G700),
    ], anchor=MIDDLE)


def overview_page(slide):
    header(slide)
    for index in range(len(ROWS)):
        row(slide, index)
    exchange(slide)
    key_band(slide)


def build(path):
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    overview_page(prs.slides.add_slide(prs.slide_layouts[6]))
    prs.save(path)


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else OUTPUT
    build(target)
    print(target)
