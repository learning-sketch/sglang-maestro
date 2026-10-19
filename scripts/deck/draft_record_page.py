"""Draft of the 30B-A3B record page: the same Trace and candidates, each verdict with its reason.

The page as it now stands in the company template keeps the white deck's colours, so the draft
is drawn in them and follows the hand-edited layout: title, scope strip, the Trace table beside
the record, the concurrency strip and one band. What changes is what the cells say -- which gate
stopped each dropped candidate and why, what the second delivery actually swapped, and a band
that reads the table instead of repeating the peak.

    python draft_record_page.py [out.pptx]
"""

import sys

from build_agent_deck import (
    CENTER,
    CW,
    DASH,
    FONT,
    MIDDLE,
    MONO,
    MX,
    RIGHT,
    SH,
    SW,
    WHITE,
    P,
    add_box,
    add_line,
    add_text,
    runs,
)
from build_model_layer_page import (
    BANDWIDTH,
    CONCURRENCY,
    DEVICE,
    PEAK,
    THEMES,
    TRACE,
    TRACED,
    UNTOUCHED,
    _bound,
    _distance,
    _ms,
    _span,
)
from pptx import Presentation
from pptx.util import Inches

T = THEMES["review"]
TITLE_BAR, TITLE_INK = "306FFF", "000000"
OUTPUT = "../../slides/record_page_draft.pptx"

TITLE = "M100 Agent：Qwen3-VL-30B-A3B 优化记录"
GROUPS = {"bound": ("大块", "读权重、读 KV"), "glue": ("小核", "搬得少、开销大")}
RECORDS = (
    # (number, change, verdict, kind, evidence lines)
    ("1", "融合 swiglu + 动态量化", "交付", "ship",
     ("逐元素 3.08 ms 与量化 0.43 ms 合成一个核：每步 3.51 → 1.75 ms",)),
    ("2", "量化 + shuffle 换成耗时恒定的核", "交付", "ship",
     ("旧核随规模从 20 涨到 141 µs；新核单测快 5.43 倍，整机吞吐 +3%",)),
    ("3", "调 grouped-GEMM 的分组旋钮", "① ③ 拦下", "drop",
     ("关掉分组算错（cos 0.000）；强制不分组慢 20%；四种模式完全同速",
      "开销在闭源核内部，模型层调参动不了，要靠算子层改核")),
    ("4", "FFN 输出合并换成 shuffle_mul_sum_sdnn", "③ 拦下", "drop",
     ("速度只有原来的 0.61 倍",)),
    ("5", "路由 softmax + topk + renorm 融合", "算过，不接", "idle",
     ("内核快 1 倍、数值也对，但只占出字间隔的 0.9%：再快也最多省 0.9%",)),
)
BAND = "Attention 搬 3.2 GB 只要 2.8 ms，小核合计只搬 0.23 GB 却要 9.5 ms；交付只换了小核"
FOOTNOTE = ("左表关 CUDA Graph 逐核计时，吞吐开 CUDA Graph 实测；专家 GEMM 的访存量取决于每步"
            "激活了哪些专家，上限未单独核算")
UNCOUNTED = "未核算"
ROW, HEAD = 0.25, 0.22


def tone(kind):
    return {"ship": T.ship, "drop": T.drop, "idle": T.g500}[kind]


def header(slide):
    add_box(slide, MX, 0.44, 0.07, 0.38, fill=TITLE_BAR)
    add_text(slide, MX + 0.22, 0.33, 10.4, 0.6, [P(TITLE, 26, TITLE_INK, True)], anchor=MIDDLE)


def scope(slide, y):
    add_box(slide, MX, y, CW, 0.34, fill=T.panel)
    add_text(slide, MX + 0.2, y, CW - 0.4, 0.34, [runs([
        ("使用模型　", 10.5, T.g900, True),
        ("Qwen3-VL-30B-A3B-Instruct-channel-fp8", 10, T.accent, True, MONO),
        ("　MoE 128 选 8 · 48 层 · M100 单卡 TP1 · 输入 / 输出各 1K", 10, T.g700),
        ("　优化目标　", 10.5, T.g900, True), ("Decode 阶段输出吞吐", 10, T.g700),
    ])], anchor=MIDDLE)


def aimed_runs(aimed):
    if aimed == "—":
        return [(aimed, 9.5, T.g400, False, "Arial")]
    kinds = {number: kind for number, _change, _verdict, kind, _lines in RECORDS}
    parts = []
    for number in aimed.split(" · "):
        if parts:
            parts.append((" · ", 9.5, T.g400, False, "Arial"))
        parts.append((number, 9.5, tone(kinds[number]), True, "Arial"))
    return parts


def trace_row(slide, x, top, scale, name, base, shipped, gb, aimed):
    bar_x, bar_y, bar_h = x + 2.00, top + (ROW - 0.13) / 2, 0.13
    length, before = _ms(shipped) * scale, _ms(base) * scale
    add_text(slide, x + 0.22, top, 1.74, ROW, [P(name, 9.5, T.g700)], anchor=MIDDLE)
    if gb is None:
        add_box(slide, bar_x, bar_y, length, bar_h, fill=T.g200, line=T.g400, dash=DASH)
    else:
        add_box(slide, bar_x, bar_y, length, bar_h, fill=T.accent)
        add_box(slide, bar_x, bar_y, max(0.02, min(_bound(gb) * scale, length)), bar_h,
                fill=T.g400)
    if before - length > 0.05:
        add_box(slide, bar_x + length, bar_y, before - length, bar_h,
                line=T.g400, line_w=0.75, dash=DASH)
    if name == UNTOUCHED:
        add_text(slide, bar_x + length + 0.08, top, 1.5, ROW,
                 [P("← 离上限最远，还没动过", 9, T.accent, True)], anchor=MIDDLE)
    distance = UNCOUNTED if gb is None else _distance(base, shipped, gb)
    for cell_x, cell_w, text in ((x + 4.12, 0.92, _span(base, shipped)),
                                 (x + 5.10, 0.82, distance)):
        quiet = text == UNCOUNTED
        add_text(slide, cell_x, top, cell_w, ROW,
                 [P(text, 9.5, T.g400 if quiet else T.g900, False, FONT if quiet else MONO,
                    align=RIGHT)], anchor=MIDDLE)
    add_text(slide, x + 5.98, top, 0.68, ROW, [runs(aimed_runs(aimed), align=CENTER)],
             anchor=MIDDLE)


def trace_panel(slide, x, y, w, h):
    add_box(slide, x, y, w, h, fill=T.panel)
    add_text(slide, x + 0.22, y + 0.06, w - 0.44, 0.28,
             [P("每类算子离理论上限多远", 12.5, T.g900, True)], anchor=MIDDLE)
    add_text(slide, x + 0.22, y + 0.33, w - 0.44, 0.22, [P(
        f"{TRACED.lstrip('c')} 并发纯 Decode · 每步耗时 ms · 理论上限 = 访存字节 ÷ "
        f"{BANDWIDTH} GB/s（实测带宽）", 9.5, T.g500)], anchor=MIDDLE)
    heads_y = y + 0.58
    add_text(slide, x + 0.22, heads_y, 1.72, 0.22, [P("类别", 9, T.g400)], anchor=MIDDLE)
    add_text(slide, x + 2.00, heads_y, 2.1, 0.22, [runs([
        ("■", 9, T.g400), (" 上限内　", 9, T.g400), ("■", 9, T.accent), (" 超出　", 9, T.g400),
        ("▭", 9, T.g400), (" 交付省掉", 9, T.g400),
    ])], anchor=MIDDLE)
    for cell_x, cell_w, align, head in ((x + 4.12, 0.92, RIGHT, "基线 → 交付"),
                                        (x + 5.10, 0.82, RIGHT, "离上限"),
                                        (x + 5.98, 0.68, CENTER, "对应改动")):
        add_text(slide, cell_x, heads_y, cell_w, 0.22,
                 [P(head, 9, T.g400, align=align)], anchor=MIDDLE)
    add_line(slide, x + 0.22, heads_y + 0.25, x + w - 0.22, heads_y + 0.25, T.g200, 0.75)

    scale = 2.05 / max(_ms(value) for row in TRACE for value in row[2:4])
    top = heads_y + 0.30
    for group, (label, gist) in GROUPS.items():
        rows = [row for row in TRACE if row[0] == group]
        base, shipped = (sum(_ms(row[index]) for row in rows) for index in (2, 3))
        share = base / float(DEVICE[0])
        add_text(slide, x + 0.22, top, w - 0.44, HEAD, [runs([
            (label, 9.5, T.accent if group == "glue" else T.g900, True),
            (f"　{gist} · {base:.2f} → {shipped:.2f} ms · 基线占每步 {share:.0%}", 9, T.g500),
        ])], anchor=MIDDLE)
        top += HEAD
        for _group, name, base_ms, shipped_ms, gb, aimed in rows:
            trace_row(slide, x, top, scale, name, base_ms, shipped_ms, gb, aimed)
            top += ROW
    add_line(slide, x + 0.22, top + 0.04, x + w - 0.22, top + 0.04, T.g200, 0.75)
    glue_saved = sum(_ms(row[2]) - _ms(row[3]) for row in TRACE if row[1] == "逐元素 + 量化")
    device_drop = 1 - float(DEVICE[1]) / float(DEVICE[0])
    add_text(slide, x + 0.22, top + 0.08, w - 0.44, HEAD, [runs([
        ("整步", 9.5, T.g900, True),
        (f"　device {' → '.join(DEVICE)} ms（−{device_drop:.0%}），其中逐元素 + 量化省 "
         f"{glue_saved:.2f} ms", 9, T.g500),
    ])], anchor=MIDDLE)


def record_panel(slide, x, y, w, h):
    add_box(slide, x, y, w, h, fill=T.panel)
    add_text(slide, x + 0.22, y + 0.08, w - 0.44, 0.28, [runs([
        ("优化记录", 12.5, T.g900, True),
        ("　10 个候选归成 5 组 · 编号对应左表最右列", 9.5, T.g500),
    ])], anchor=MIDDLE)
    verdict_w = 1.1
    top = y + 0.46
    for number, change, verdict, kind, lines in RECORDS:
        entry_h = 0.30 + 0.19 * len(lines)
        shipped = kind == "ship"
        if shipped:
            add_box(slide, x + 0.14, top, w - 0.28, entry_h, fill=T.tint)
        add_text(slide, x + 0.24, top + 0.04, 0.28, 0.24,
                 [P(number, 10.5, tone(kind), True, "Arial")], anchor=MIDDLE)
        add_text(slide, x + 0.56, top + 0.04, w - 0.80 - verdict_w, 0.24,
                 [P(change, 10, T.g900, shipped)], anchor=MIDDLE)
        add_text(slide, x + w - 0.24 - verdict_w, top + 0.04, verdict_w, 0.24,
                 [P(verdict, 10, tone(kind), True, align=RIGHT)], anchor=MIDDLE)
        for index, line in enumerate(lines):
            add_text(slide, x + 0.56, top + 0.29 + index * 0.19, w - 0.80, 0.19,
                     [P(line, 9, T.g700)], anchor=MIDDLE)
        top += entry_h + 0.06


def concurrency(slide, y, h):
    lead_w = 3.05
    add_text(slide, MX, y, lead_w, h, [
        P("同一交付，换个并发", 11.5, T.g900, True, after=2),
        P("输出吞吐 tok/s，并发越高收益越大：", 9.5, T.g500),
        P("换掉的小核随规模变慢（第 2 组 20 → 141 µs）", 9.5, T.g500),
    ], anchor=MIDDLE)
    gap = 0.14
    x0 = MX + lead_w + 0.1
    cell_w = (SW - MX - x0 - (len(CONCURRENCY) - 1) * gap) / len(CONCURRENCY)
    for index, (label, base, shipped) in enumerate(CONCURRENCY):
        cx = x0 + index * (cell_w + gap)
        hero = label == PEAK
        tag = {TRACED: f"{label} · 左表工作点", PEAK: f"{label} · 峰值"}.get(label, label)
        add_box(slide, cx, y, cell_w, h, fill=T.tint if hero else WHITE,
                line=None if hero else T.g200)
        add_text(slide, cx + 0.16, y + 0.04, cell_w - 0.32, 0.19,
                 [P(tag, 9.5, T.accent if hero else T.g500, hero)], anchor=MIDDLE)
        add_text(slide, cx + 0.16, y + 0.23, cell_w - 0.32, 0.22,
                 [P(f"+{shipped / base - 1:.1%}", 15, T.g900, True, "Arial")], anchor=MIDDLE)
        add_text(slide, cx + 0.16, y + 0.45, cell_w - 0.32, 0.16,
                 [P(f"{base:.1f} → {shipped:.1f}", 9, T.g500, False, MONO)], anchor=MIDDLE)


def record_page(slide):
    header(slide)
    scope(slide, 1.12)
    panel_y, panel_h, trace_w = 1.58, 3.38, 6.88
    trace_panel(slide, MX, panel_y, trace_w, panel_h)
    list_x = MX + trace_w + 0.2
    record_panel(slide, list_x, panel_y, SW - MX - list_x, panel_h)
    concurrency(slide, 5.10, 0.64)
    T.band(slide, 5.88, 0.54, "结论", BAND)
    add_text(slide, MX, 6.50, CW, 0.26, [runs([
        ("口径　", 10, T.accent, True), (FOOTNOTE, 10, T.g500),
    ])], anchor=MIDDLE)


def build(path):
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    record_page(prs.slides.add_slide(prs.slide_layouts[6]))
    prs.save(path)


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else OUTPUT
    build(target)
    print(target)
